"""Columns inserted into an opened workbook by a write-back are moved in the saved file: cells with their
styles, column widths and every formula in the workbook re-referenced, or the save is refused where
the prototype cannot move a feature. Run with the workbook driver and openpyxl."""
import json, subprocess, sys, tempfile
from pathlib import Path
import openpyxl
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'Tests'))

driver = subprocess.Popen([sys.argv[1], str(ROOT / 'FullEngine/publish/StatsDirectEngine.dylib')], stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)
def request(**payload):
    driver.stdin.write(json.dumps(payload) + '\n'); driver.stdin.flush()
    return json.loads(driver.stdout.readline())
def cell(col, row, text, kind='number'): return {'col': col, 'row': row, 'text': str(text), 'kind': kind}

try:
    with tempfile.TemporaryDirectory() as folder:
        source = Path(folder) / 'source.xlsx'
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Data'
        ws['A1'] = 'double'; ws['B1'] = 'y'; ws['C1'] = 'z'
        for r, (y, z) in enumerate([(0.5, 7), (0.25, 8), (0.125, 9)], start=2):
            ws.cell(r, 1, f'=B{r}*2'); c = ws.cell(r, 2, y); c.number_format = '0.00%'; ws.cell(r, 3, z)
        ws.column_dimensions['B'].width = 30
        other = wb.create_sheet('Other'); other['A1'] = '=Data!B2+1'
        wb.save(source)
        opened = request(action='open', path=str(source)); assert 'error' not in opened, opened
        # The grid inserted one column before B (a log transform written after the selected column A):
        # B holds the output, y moved to C with its format, z to D; A's formulas and Other's formula still point at y.
        moved = [cell(2, r, y) for r, y in enumerate([0.5, 0.25, 0.125], start=1)] + [cell(3, r, z) for r, z in enumerate([7, 8, 9], start=1)] + [cell(2, 0, 'y', 'text'), cell(3, 0, 'z', 'text')]
        output = [cell(1, 0, 'log', 'text')] + [cell(1, r, v) for r, v in enumerate([10, 20, 30], start=1)]
        saved = Path(folder) / 'inserted.xlsx'
        result = request(action='save', id=opened['id'], path=str(saved), sheets=[{'name': 'Data', 'cells': output + moved}], inserts=[{'name': 'Data', 'inserts': [{'col': 1, 'count': 1}]}])
        assert result.get('ok'), result
        book = openpyxl.load_workbook(saved); data = book['Data']
        assert [data.cell(1, c).value for c in range(1, 5)] == ['double', 'log', 'y', 'z'], [data.cell(1, c).value for c in range(1, 5)]
        assert [data.cell(r, 2).value for r in (2, 3, 4)] == [10, 20, 30]
        assert [data.cell(r, 3).value for r in (2, 3, 4)] == [0.5, 0.25, 0.125] and data['C2'].number_format == '0.00%', (data['C2'].value, data['C2'].number_format)
        assert [data.cell(r, 4).value for r in (2, 3, 4)] == [7, 8, 9]
        assert data['A2'].value == '=C2*2' and data['A4'].value == '=C4*2', (data['A2'].value, data['A4'].value)   # re-referenced past the inserted column
        assert book['Other']['A1'].value == '=Data!C2+1', book['Other']['A1'].value
        assert abs(data.column_dimensions['C'].width - 30) < 0.01, data.column_dimensions['C'].width   # the width travelled with y
        assert data.max_column == 4
        print('PASS: an inserted column moves the file\'s cells with their styles and widths and re-references every formula in the workbook')

        # Two insertions in sequence, the second in the coordinates after the first.
        twice = Path(folder) / 'twice.xlsx'
        result = request(action='save', id=opened['id'], path=str(twice), sheets=[{'name': 'Data', 'cells': [cell(1, 0, 'first', 'text'), cell(3, 0, 'second', 'text'), cell(2, 0, 'y', 'text'), cell(4, 0, 'z', 'text')] + [cell(2, r, y) for r, y in enumerate([0.5, 0.25, 0.125], start=1)] + [cell(4, r, z) for r, z in enumerate([7, 8, 9], start=1)]}],
                         inserts=[{'name': 'Data', 'inserts': [{'col': 1, 'count': 1}, {'col': 3, 'count': 1}]}])
        assert result.get('ok'), result
        data = openpyxl.load_workbook(twice)['Data']
        assert [data.cell(1, c).value for c in range(1, 6)] == ['double', 'first', 'y', 'second', 'z'] and data['A2'].value == '=C2*2' and data['E3'].value == 8, [data.cell(1, c).value for c in range(1, 6)]
        print('PASS: successive insertions are applied in order')

        import shutil
        node=sys.argv[2] if len(sys.argv)>2 else shutil.which('node')
        assert node, 'Supply Node to test the actual grid export'
        grid=json.loads(subprocess.check_output([node,str(ROOT/'Tests/insert-workbook.mjs')],input=json.dumps(opened),text=True))
        for state in ['applied','undone','redone']:
            dest=Path(folder)/(state+'.xlsx')
            for repeat in range(2):
                result=request(action='save',id=opened['id'],path=str(dest),**grid[state]); assert result.get('ok'),result
            x=openpyxl.load_workbook(dest)['Data']
            if state=='undone': assert x['B3'].value==21 and x['A3'].value=='=B3*2'
            else:
                assert x['C3'].value==21 and x['B3'].value=='=C3*2' and x['D3'].value==8
                assert x['C3'].number_format=='0.00%' and openpyxl.load_workbook(dest,data_only=True)['Data']['B3'].value==42
                assert len(grid[state]['sheets'][0]['cells'])==5,grid[state]
        print('PASS: real grid plan → export → native save, prior edits, moved formula cells, recalculation, repeated saves and undo/redo')

        # Every refusal is checked at preflight as well as save; neither the source nor
        # an existing destination may change, and no temporary file may remain.
        import hashlib, zipfile, xml.etree.ElementTree as ET
        from openpyxl.styles import Font, PatternFill, Border, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.workbook.defined_name import DefinedName
        from openpyxl.worksheet.table import Table
        from openpyxl.worksheet.formula import ArrayFormula
        NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        def insert(path, col=1, count=1, stream=False, expect=None, sheet='Data', changes=None):
            opened=request(action='open',path=str(path)); assert 'error' not in opened, opened
            args=dict(id=opened['id'],inserts=[{'name':sheet,'inserts':[{'col':col,'count':count}]}])
            digest=hashlib.sha256(path.read_bytes()).digest()
            pre=request(action='validateInserts',**args)
            dest=Path(folder)/('result-'+path.name)
            dest.write_bytes(b'previous destination')
            result=request(action='save',path=str(dest),stream=stream,sheets=[{'name':sheet,'cells':changes or [cell(col,0,'result','text')]}],**args)
            if expect:
                assert expect.lower() in pre.get('error','').lower(), pre
                assert expect.lower() in result.get('error','').lower(), result
                assert dest.read_bytes()==b'previous destination'
            else:
                assert pre.get('ok') and result.get('ok'), (pre,result)
                if stream and opened.get('formulaCount',0): assert result.get('uncachedFormulas',0)>0,result
            assert hashlib.sha256(path.read_bytes()).digest()==digest
            assert not list(Path(folder).glob('*.tmp*'))
            request(action='close',id=opened['id'])
            return dest
        def book(name):
            w=openpyxl.Workbook(); d=w.active; d.title='Data'
            d.append(['a','b','c']); d.append([1,2,3]);
            return w,d,Path(folder)/name
        w,d,p=book('features.xlsx')
        d['B2'].font=Font(name='Arial',bold=True,color='FF0000')
        d['B2'].fill=PatternFill('solid',fgColor='00FF00'); d['B2'].number_format='0.00%'
        d['B2'].border=Border(bottom=Side(style='thin')); d.column_dimensions['B'].width=32
        d.column_dimensions['C'].hidden=True; d.row_dimensions[2].height=30
        d.merge_cells('D5:E5'); d['D5']='merged'
        d.conditional_formatting.add('B2:C6',FormulaRule(formula=['B2>0'],fill=PatternFill('solid',fgColor='FFFF00')))
        dv=DataValidation(type='whole',operator='between',formula1='B10',formula2='C10'); dv.add('B2:C6'); d.add_data_validation(dv)
        d['B3']='link'; d['B3'].hyperlink='https://statsdirect.com'; d.freeze_panes='B2'
        w.defined_names.add(DefinedName('Results',attr_text="'Data'!$B$2:$C$6")); d.print_area='A1:C6'
        d.auto_filter.ref='A1:C6'; d.auto_filter.add_filter_column(1,['2'])
        d['C8']='=SUM($B2,C$2,$B$2,B:B,2:2)+N("B2")'
        other=w.create_sheet('Other'); other['A1']="='Data'!B2+1"; other['A2']='=[1]Data!B2'
        w.save(p)
        for streaming in (False,True):
            saved=insert(p,stream=streaming); loaded=openpyxl.load_workbook(saved); x=loaded['Data']
            assert x['C2'].value==2 and x['C2'].number_format=='0.00%' and x['C2'].font.bold and x['C2'].font.name=='Arial'
            assert x['C2'].fill.fgColor.rgb=='0000FF00' and x['C2'].border.bottom.style=='thin'
            assert x.column_dimensions['C'].width==32 and x.column_dimensions['D'].hidden and x.row_dimensions[2].height==30
            assert str(x.merged_cells)=='E5:F5' and x['E5'].value=='merged'
            cf=list(x.conditional_formatting)[0]; assert str(cf.sqref)=='C2:D6'; assert x.conditional_formatting[cf][0].formula==['C2>0']
            validation=x.data_validations.dataValidation[0]; assert str(validation.sqref)=='C2:D6' and validation.formula1=='C10' and validation.formula2=='D10'
            assert x['C3'].hyperlink.target=='https://statsdirect.com' and x['C3'].hyperlink.ref=='C3'
            assert x.freeze_panes=='C2' and x.sheet_view.pane.xSplit==1
            assert loaded.defined_names['Results'].attr_text.replace("'",'')=='Data!$C$2:$D$6'
            assert '$A$1:$D$6' in str(x.print_area)
            assert x.auto_filter.ref=='A1:D6' and x.auto_filter.filterColumn[0].colId==2
            assert x['D8'].value=='=SUM($C2,D$2,$C$2,C:C,2:2)+N("B2")', x['D8'].value
            assert loaded['Other']['A1'].value.replace("'",'')=='=Data!C2+1'
            assert loaded['Other']['A2'].value=='=[1]Data!B2'
        print('PASS: formatting, widths, hidden columns, merged ranges, conditional formats, validation, links, freeze panes, filters, names and cross-sheet/absolute/external references on both save paths')

        w,d,p=book('table.xlsx'); d.add_table(Table(displayName='Trial',ref='A1:C2')); w.save(p)
        saved=insert(p,col=0); assert openpyxl.load_workbook(saved)['Data'].tables['Trial'].ref=='B1:D2'
        expanded=openpyxl.load_workbook(insert(p,col=1))['Data']; assert expanded.tables['Trial'].ref=='A1:D2' and [c.name for c in expanded.tables['Trial'].tableColumns]==['a','result','b','c']
        w,d,p=book('merged.xlsx'); d.merge_cells('A3:B3'); w.save(p)
        assert str(openpyxl.load_workbook(insert(p)).active.merged_cells)=='A3:C3'
        w,d,p=book('array.xlsx'); d['B4']=ArrayFormula(ref='B4:C4',text='=A2:B2*2'); w.save(p)
        for streaming in (False,True):
            saved=insert(p,stream=streaming); a=openpyxl.load_workbook(saved)['Data']['C4'].value
            assert a.ref=='C4:D4' and a.text=='=A2:C2*2', (a.ref,a.text)
        insert(p,col=2,expect='split an array formula')
        w,d,p=book('protected.xlsx'); d.protection.sheet=True; w.save(p); insert(p,expect='protected')
        d.protection.insertColumns=False; w.save(p); insert(p)
        w,d,p=book('edge.xlsx'); d.cell(2,16384,'keep'); w.save(p); insert(p,expect='last column')
        w,d,p=book('3d.xlsx'); w.create_sheet('Other'); d['A3']='=SUM(Data:Other!B2)'; w.save(p); assert openpyxl.load_workbook(insert(p))['Data']['A3'].value=='=SUM(Data:Other!B2)'
        print('PASS: whole tables and arrays move; splitting arrays, protected insertions, and overflow are refused atomically; divergent 3-D references retain Excel semantics')

        # The shipped example workbook uses shared formulas. Expand the expressions,
        # then structurally adjust each one rather than discarding sharing metadata only.
        saved=insert(ROOT/'Content/Examples/test.xlsx',sheet='Parametric')
        source_book=openpyxl.load_workbook(ROOT/'Content/Examples/test.xlsx')
        target_book=openpyxl.load_workbook(saved)
        assert target_book['Parametric'].max_column==source_book['Parametric'].max_column+1
        assert target_book['Parametric']['C2'].value==source_book['Parametric']['B2'].value
        with zipfile.ZipFile(saved) as z:
            assert all(b't="shared"' not in z.read(n) for n in z.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'))
        print('PASS: insertion into the shipped StatsDirect example workbook with shared formulas')

        from openpyxl.chart import BarChart, Reference
        w,d,p=book('chart.xlsx'); chart=BarChart(); chart.add_data(Reference(d,min_col=2,min_row=1,max_row=2),titles_from_data=True); d.add_chart(chart,'E2'); w.save(p)
        saved=insert(p)
        moved_chart=openpyxl.load_workbook(saved)['Data']._charts[0]
        assert moved_chart.anchor._from.col==5 and moved_chart.series[0].val.numRef.f.replace("'",'')=='Data!$C$2'
        with zipfile.ZipFile(p) as original, zipfile.ZipFile(saved) as moved:
            assert original.read('xl/styles.xml')==moved.read('xl/styles.xml')
        from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
        for mode in ['twoCell','oneCell','absolute']:
            for position in [0,1]:
                w,d,p=book('anchor-'+mode+str(position)+'.xlsx'); chart=BarChart(); chart.add_data(Reference(d,min_col=2,min_row=1,max_row=2))
                chart.anchor=TwoCellAnchor(editAs=mode,_from=AnchorMarker(col=0,row=0),to=AnchorMarker(col=3,row=4)); d.add_chart(chart); w.save(p)
                a=openpyxl.load_workbook(insert(p,col=position))['Data']._charts[0].anchor
                expected=(0,3) if mode=='absolute' or mode=='oneCell' and position==1 else (1,4) if position==0 else (0,4)
                assert (a._from.col,a.to.col)==expected,(mode,position,a._from.col,a.to.col)
        from openpyxl.comments import Comment
        w,d,p=book('comments.xlsx'); d['B2'].comment=Comment('Keep this comment','Tester'); w.save(p)
        for streaming in (False,True):
            moved=openpyxl.load_workbook(insert(p,stream=streaming))['Data']['C2'].comment
            assert moved.text=='Keep this comment' and moved.author=='Tester'
        # Excel-authored notes have an eight-coordinate anchor. Exercise move/size
        # options separately, including explicit False (presence alone is not true).
        for mode,flags,expected in [
            ('sized','<x:MoveWithCells/><x:SizeWithCells/>',(0,4)),
            ('move','<x:MoveWithCells/><x:SizeWithCells>False</x:SizeWithCells>',(0,3)),
            ('fixed','<x:MoveWithCells>False</x:MoveWithCells>',(0,3))]:
            anchored=Path(folder)/('anchored-'+mode+'.xlsx')
            with zipfile.ZipFile(p) as z,zipfile.ZipFile(anchored,'w') as out:
                for item in z.infolist():
                    data=z.read(item.filename)
                    if item.filename.endswith('.vml'):
                        xml=ET.fromstring(data);x='urn:schemas-microsoft-com:office:excel'
                        client=xml.find('.//{'+x+'}ClientData')
                        for name in ['MoveWithCells','SizeWithCells','Anchor']:
                            for el in list(client.findall('{'+x+'}'+name)):client.remove(el)
                        extra=ET.fromstring('<extra xmlns:x="'+x+'">'+flags+'<x:Anchor>0, 6, 0, 2, 3, 8, 4, 1</x:Anchor></extra>')
                        client.extend(list(extra));data=ET.tostring(xml)
                    out.writestr(item,data)
            for streaming in (False,True):
                saved=insert(anchored,stream=streaming)
                with zipfile.ZipFile(saved) as z:
                    vml=ET.fromstring(z.read(next(n for n in z.namelist() if n.endswith('.vml'))))
                    anchor=list(map(int,vml.find('.//{'+x+'}Anchor').text.split(',')))
                    assert (anchor[0],anchor[4])==expected,(mode,streaming,anchor)
                    assert vml.find('.//{'+x+'}Column').text=='2'
        # A legacy form control shares the VML format and is now preserved.
        control=Path(folder)/'control.xlsx'
        with zipfile.ZipFile(p) as z,zipfile.ZipFile(control,'w') as out:
            for item in z.infolist():
                data=z.read(item.filename)
                if item.filename.endswith('.vml'): data=data.replace(b'ObjectType="Note"',b'ObjectType="Button"')
                out.writestr(item,data)
        insert(control)
        w,d,p=book('ambiguous-name.xlsx'); w.defined_names.add(DefinedName('Relative',attr_text='$B$2')); w.save(p); insert(p,expect='unqualified cell reference')
        print('PASS: chart anchor/series and legacy notes move, original style XML remains identical, controls are preserved and ambiguous names are refused before mutation')

        # Truly cross the million-cell threshold; streaming is selected automatically.
        large=Path(folder)/'large.xlsx'; w=openpyxl.Workbook(write_only=True); d=w.create_sheet('Data')
        for r in range(100001): d.append([r]*10)
        w.save(large)
        saved=insert(large,col=1)
        read=openpyxl.load_workbook(saved,read_only=True); rows=read['Data'].iter_rows(values_only=True)
        first=next(rows); assert first[:4]==(0,'result',0,0),first
        last=None
        for last in rows: pass
        assert last==(100000,None,*([100000]*9)),last
        read.close()
        print('PASS: automatic streaming insertion moves 1,000,010 cells without a size refusal or data loss')
finally:
    driver.stdin.close(); driver.wait()
