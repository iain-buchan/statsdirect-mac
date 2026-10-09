"""Read Excel formats through the real native bridge; independently check exported data.

python3 Tests/test_excel_compatibility.py <workbook-driver> <node>
"""
import datetime
import hashlib
import json
import re
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from zipfile import ZipFile, ZIP_DEFLATED
import openpyxl
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'Tests/Fixtures/Excel'
driver = subprocess.Popen([sys.argv[1], str(ROOT / 'FullEngine/publish/StatsDirectEngine.dylib')],
                          stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)

def request(**payload):
    driver.stdin.write(json.dumps(payload) + '\n'); driver.stdin.flush()
    return json.loads(driver.stdout.readline())

def opened(path, **kwargs):
    reply = request(action='open', path=str(path), **kwargs)
    assert 'error' not in reply, (str(path), reply)
    return reply

def close(reply):
    assert request(action='close', id=reply['id'])['ok']

def values(reply, sheet=0):
    return {(c['col'], c['row']): (c['text'], c['kind']) for c in reply['sheets'][sheet]['cells']}

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def independent_values(path):
    result = {}
    for row in openpyxl.load_workbook(path, data_only=True).active:
        for c in row:
            if c.value is not None: result[(c.column - 1, c.row - 1)] = c.value
    return result

def exported(reply):
    # Exercise the grid model: data copies have an ID for lifetime/save protection,
    # but must export ALL cells, not just changes as a preserving XLSX edit does.
    script = "import fs from 'node:fs';import {WorkbookStore} from './Grid/workbook.mjs';const b=new WorkbookStore();b.load(JSON.parse(fs.readFileSync(0,'utf8')));process.stdout.write(JSON.stringify(b.export()));"
    return json.loads(subprocess.check_output([sys.argv[2], '--input-type=module', '-e', script],
        input=json.dumps(reply), text=True, cwd=ROOT))

def variant(source, target, content_type=None, transform=None):
    with ZipFile(source) as src, ZipFile(target, 'w', ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == '[Content_Types].xml' and content_type:
                data = data.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml', content_type.encode())
            if transform: data = transform(item.filename, data)
            dst.writestr(item, data)

try:
    with tempfile.TemporaryDirectory() as folder:
        folder = Path(folder)
        expected = independent_values(FIXTURES / '10x10.xlsx')
        for name in ['10x10.xls', '10x10.xlsb', 'strict/10x10.xlsx']:
            source = FIXTURES / name; before = digest(source)
            reply = opened(source)
            assert reply['dataCopy'] and reply['sheets'][0]['rows'] == 10 and reply['sheets'][0]['columns'] == 10
            actual = {k: float(v) if kind == 'number' else v for k, (v, kind) in values(reply).items()}
            assert actual == expected, (name, actual)
            out = folder / (source.stem + '-' + source.suffix[1:] + '-data.xlsx')
            saved = request(action='save', id=reply['id'], path=str(out), **exported(reply))
            assert saved.get('ok'), saved
            assert independent_values(out) == expected
            assert digest(source) == before
            close(reply)
        print('PASS: XLS, XLSB and Strict XLSX import every cell at its original address, export through the grid, and match an independent reader')

        legacy = opened(FIXTURES / 'EncodingFormulaDate1520.xls')
        binary = opened(FIXTURES / 'EncodingFormulaDate1520.xlsb')
        assert len(legacy['sheets']) == len(binary['sheets']) == 3
        # Upstream's XLSB fixture has a deliberately different B6 date.
        assert values(legacy)[(1, 5)] == ('2010-09-09 00:00:00', 'datetime')
        assert values(binary)[(1, 5)] == ('2019-11-01 00:00:00', 'datetime')
        assert {k:v for k,v in values(legacy).items() if k != (1,5)} == {k:v for k,v in values(binary).items() if k != (1,5)}
        assert values(legacy)[(1, 1)] == ('2009-05-01 00:00:00', 'datetime')
        assert values(legacy)[(4, 2)][0].endswith('11:00:00')
        for reply in [legacy, binary]: close(reply)
        for fmt in [2, 3, 4, 5]:
            reply = opened(FIXTURES / f'as3xls_BIFF{fmt}.xls')
            cells = values(reply)
            assert cells[(0, 0)] == ('1', 'number') and cells[(1, 0)] == ('Hi', 'text')
            assert cells[(2, 1)] == ('2007-02-22 00:00:00', 'datetime')
            close(reply)
        for ext in ['xls', 'xlsb']:
            reply = opened(FIXTURES / f'Issue329_Error.{ext}')
            assert [v for v in values(reply).values()] == [('#DIV/0!', 'error'), ('#N/A', 'error'), ('#VALUE!', 'error')]
            close(reply)
        print('PASS: older BIFF2–5, cached results, dates, times, text and Excel errors retain their types; empty worksheets remain accessible')

        source = FIXTURES / 'AllColumnsNotReadInHiddenTable.xls'
        reply = opened(source, snapshot=str(folder / 'out-of-order'))
        assert reply['sheets'][1]['rows'] == 51
        decoded = json.loads(subprocess.check_output([sys.argv[2], str(ROOT / 'Tests/decode-snapshot.mjs'), reply['sheets'][1]['snapshot']], text=True))
        row = sorted((c for c in decoded[0]['cells'] if c['row'] == 1), key=lambda c:c['col'])
        assert row[0]['text'] == '21/09/2015'
        assert [float(c['text']) for c in row[1:]] == [1187.5282349881188,650.8582749049624,1361.7209439645526,321.74647548613916,369.48879457369037]
        assert len(list((folder / 'out-of-order').glob('*.sdcol'))) == len(reply['sheets'])
        close(reply)
        reply = opened(FIXTURES / 'Issue14_InvalidOADate.xlsx')
        assert values(reply)[(0, 0)] == ('1000000000000', 'number')
        close(reply)
        print('PASS: out-of-order legacy rows are reconstructed without dropped columns or partial snapshots; out-of-range date-formatted numbers remain readable')

        for name in ['Issue242_StdRc4PwdPassword.xls', 'Issue242_XorPwdPassword.xls',
                     'agile_AES256_SHA512_CBC_pwd_password.xlsx', 'standard_AES128_SHA1_ECB_pwd_password.xlsx',
                     'agile_AES128_SHA1_CBC_pwd_password.xlsb']:
            source = folder / name; shutil.copy2(FIXTURES / name, source); before = digest(source)
            snapshots = folder / 'password-snapshots'
            for password in [None, 'wrong-password']:
                failed = request(action='open', path=str(source), snapshot=str(snapshots), password=password)
                assert failed.get('code') == 'workbookPassword', failed
                assert not list(snapshots.glob('*'))
            reply = opened(source, password='password', snapshot=str(snapshots))
            assert reply['dataCopy'] and reply['sheets'][0]['cells'] == []
            snap = Path(reply['sheets'][0]['snapshot'])
            decoded = json.loads(subprocess.check_output([sys.argv[2], str(ROOT / 'Tests/decode-snapshot.mjs'), str(snap)], text=True))
            assert decoded[0]['cells'][0]['text'] == 'Password: password'
            assert not list(snapshots.glob('statsdirect-workbook-*'))  # no decrypted source retained
            if source.suffix == '.xlsx':
                saved = request(action='save', id=reply['id'], path=str(source), sheets=decoded)
                assert 'new name' in saved.get('error', ''), saved
            out = folder / (source.name + '-data.xlsx')
            assert request(action='save', id=reply['id'], path=str(out), snapshot=str(snap)).get('ok')
            assert openpyxl.load_workbook(out).active['A1'].value == 'Password: password'
            assert digest(source) == before
            close(reply); snap.unlink()
        print('PASS: encrypted XLS/XLSX/XLSB prompt, reject wrong passwords, read with the correct password, export values, and leave originals unchanged')

        # The features that prevent safe column insertion must NEVER prevent reading.
        rich = folder / 'complex.xlsx'
        book = openpyxl.Workbook(); ws = book.active; ws.title = 'Data'
        for row in [['dose', 'response', 'cached'], [1, 12.5, '=A2+B2'], [2, 17.5, '=A3+B3']]: ws.append(row)
        ws.add_table(Table(displayName='Measures', ref='A1:C3'))
        ws.merge_cells('E1:F1'); ws['E1'] = 'Merged heading'
        ws.add_data_validation(DataValidation(type='whole', formula1=0, formula2=100))
        ws.data_validations.dataValidation[0].add('A2:A3')
        ws.protection.sheet = True
        book.defined_names.add(DefinedName('Dose', attr_text="'Data'!$A$2:$A$3"))
        hidden = book.create_sheet('Hidden data'); hidden.sheet_state = 'veryHidden'; hidden['D7'] = '00123'
        types = book.create_sheet('Types'); types.append([True, datetime.datetime(2024, 2, 29, 12, 30), 'München', '00123'])
        book.save(rich)
        cached = folder / 'cached.xlsx'
        def cache_formulas(part, data):
            if part == 'xl/worksheets/sheet1.xml':
                for row, value in [(2, '13.5'), (3, '19.5')]:
                    formula = f'<f>A{row}+B{row}</f>'.encode()
                    data = re.sub(re.escape(formula) + rb'<v(?:\s*/>|></v>)', formula + f'<v>{value}</v>'.encode(), data)
            return data
        variant(rich, cached, transform=cache_formulas)
        variants = [('xlsx', None), ('xlsm', 'application/vnd.ms-excel.sheet.macroEnabled.main+xml'),
                    ('xltx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml'),
                    ('xltm', 'application/vnd.ms-excel.template.macroEnabled.main+xml')]
        for ext, content_type in variants:
            source = folder / f'complex.{ext}'; variant(cached, source, content_type)
            before = digest(source); reply = opened(source)
            assert [s['name'] for s in reply['sheets']] == ['Data', 'Hidden data', 'Types']
            assert reply['sheets'][1]['hidden'] and values(reply, 1)[(3, 6)] == ('00123', 'text'), (ext, reply['sheets'][1])
            assert values(reply)[(2, 1)] == ('13.5', 'number'), (ext, values(reply)[(2, 1)])
            assert values(reply, 2)[(0, 0)] == ('TRUE', 'boolean')
            assert values(reply, 2)[(1, 0)] == ('2024-02-29 12:30:00', 'datetime')
            assert values(reply, 2)[(2, 0)] == ('München', 'text')
            assert bool(reply.get('dataCopy')) == (ext != 'xlsx')
            typed = opened(source, snapshot=str(folder / 'typed'))
            for plain, sheet in zip(reply['sheets'], typed['sheets']):
                decoded = json.loads(subprocess.check_output([sys.argv[2], str(ROOT / 'Tests/decode-snapshot.mjs'), sheet['snapshot']], text=True))
                assert decoded[0]['cells'] == plain['cells']
                Path(sheet['snapshot']).unlink()
            close(typed); close(reply); assert digest(source) == before
        template = folder / 'legacy.xlt'; shutil.copy2(FIXTURES / '10x10.xls', template)
        reply = opened(template); assert reply['cellCount'] == len(expected); close(reply)
        print('PASS: macro-enabled and template workbooks, protected sheets, tables, merged cells, validation, named ranges and hidden sheets all open; typed snapshots match cell data')

        # Exceed the former million-cell editing threshold while reading, and reach
        # Excel's last row and column. Only populated cells consume snapshot storage.
        large = folder / 'full-height.xlsm'
        with ZipFile(FIXTURES / '10x10.xlsx') as src, ZipFile(large, 'w', ZIP_DEFLATED) as dst:
            for item in src.infolist():
                if item.filename != 'xl/worksheets/sheet1.xml': dst.writestr(item, src.read(item.filename))
            with dst.open('xl/worksheets/sheet1.xml', 'w') as sheet:
                sheet.write(b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>')
                for row in range(1, 1048577):
                    tail = '<c r="XFD1048576"><v>42</v></c>' if row == 1048576 else ''
                    sheet.write(f'<row r="{row}"><c r="A{row}"><v>{row}</v></c>{tail}</row>'.encode())
                sheet.write(b'</sheetData></worksheet>')
        big = opened(large, snapshot=str(folder / 'large-snapshots'))
        assert big['cellCount'] == 1048577
        assert (big['sheets'][0]['rows'], big['sheets'][0]['columns']) == (1048576, 16384)
        decoded = json.loads(subprocess.check_output([sys.argv[2], '--input-type=module', '-e',
            "import fs from 'node:fs';import {decodeSnapshot} from './Grid/snapshot.mjs';const b=fs.readFileSync(process.argv[1]);const s=decodeSnapshot(b.buffer.slice(b.byteOffset,b.byteOffset+b.byteLength)).sheets[0];console.log(JSON.stringify(s.columns.map(c=>({col:c.col,count:c.rows.length,first:c.nums[0],last:c.nums[c.nums.length-1],lastRow:c.rows[c.rows.length-1]}))));",
            big['sheets'][0]['snapshot']], text=True, cwd=ROOT))
        assert decoded == [{'col':0,'count':1048576,'first':1,'last':1048576,'lastRow':1048575},
                           {'col':16383,'count':1,'first':42,'last':42,'lastRow':1048575}], decoded
        close(big)
        print('PASS: 1,048,577 populated cells import without the editing limit, reaching row 1,048,576 and column XFD with correct snapshot values')

        # Format detection uses the file, not a possibly misleading extension.
        renamed = folder / 'old-workbook.xlsx'; shutil.copy2(FIXTURES / '10x10.xls', renamed)
        reply = opened(renamed); assert reply['dataCopy'] and reply['cellCount'] == len(expected); close(reply)
        macro_renamed = folder / 'macro-renamed.xlsx'; shutil.copy2(folder / 'complex.xlsm', macro_renamed)
        reply = opened(macro_renamed); assert reply['dataCopy']; close(reply)
        # A malformed later worksheet must clean up the first sheet's snapshot.
        broken = folder / 'broken.xlsm'
        variant(folder / 'complex.xlsm', broken, transform=lambda part, data: b'<broken' if part == 'xl/worksheets/sheet2.xml' else data)
        snaps = folder / 'failed'
        assert 'error' in request(action='open', path=str(broken), snapshot=str(snaps))
        assert not list(snaps.glob('*'))
        for ext in ['xls', 'xlsb', 'xlsm', 'xlsx']:
            bad = folder / ('corrupt.' + ext); bad.write_text('not an Excel workbook')
            failed = request(action='open', path=str(bad), snapshot=str(snaps))
            assert 'error' in failed and failed.get('code') != 'workbookPassword'
        print('PASS: mislabeled files are detected; corrupt files fail without partial worksheets, leftover snapshots or spurious password prompts')
finally:
    driver.stdin.close(); driver.wait(timeout=15)
