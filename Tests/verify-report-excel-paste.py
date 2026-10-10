"""Inspect the workbook saved by Excel after pasting the presentation fixture.

Copy the presentation-report.html contents in the native report editor, paste
into A1 of a blank Excel sheet, save as .xlsx, then pass that file here.
"""
from openpyxl import load_workbook
import sys
s=load_workbook(sys.argv[1]).active
for address,value,fmt in [('B4',12.5,'0.000'),('B5',.95,'0.00%'),('B6',1.2e-7,'0.00E+00'),('B7',-2.5,'0.00')]:
    c=s[address]
    assert (c.value,c.data_type,c.number_format)==(value,'n',fmt),(address,c.value,c.number_format)
    assert c.alignment.horizontal=='left'
for address,value in [('B8','0012'),('B9','1234567890123456'),('B10','=1+1')]:
    c=s[address]
    assert (c.value,c.data_type,c.number_format)==(value,'s','@'),address
for row in s.iter_rows(min_row=3,max_row=10,max_col=2):
    for c in row:
        assert (c.font.name,c.font.sz,c.font.b)==('Arial',10,False),c.coordinate
        assert all(getattr(c.border,edge).style is None for edge in ['top','bottom','left','right'])
assert s['A3'].font.u=='single' and s['B3'].font.u=='single'
assert (s['A1'].font.name,s['A1'].font.sz,s['A1'].font.b)==('Georgia',16,False)
c=s['A11']
assert (c.font.name,c.font.sz,c.font.color.rgb)==('Georgia',18,'FF663399')
assert c.fill.fgColor.rgb=='FFFFFF00' and c.alignment.horizontal=='center'
assert c.border.top.color.rgb=='FFCC0000' and c.border.bottom.color.rgb=='FF0000CC'
assert c.border.bottom.style in {'dashed','mediumDashed'}
assert s['B11'].fill.fgColor.rgb=='FFDDFFEE'
assert (s['A12'].font.name,s['A12'].font.sz)==('Georgia',14)
print('PASS: native Excel save retains fonts, underlines, border colours/styles, fills, real numeric values/formats and text identifiers/formula-like labels')
