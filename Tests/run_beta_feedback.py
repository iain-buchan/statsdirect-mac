"""Native macOS integration checks. Add --live to test the signed-in ChatGPT account with bundled example data."""
from pathlib import Path
import argparse,subprocess,plistlib,uuid
root=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--app',type=Path,default=root/'StatsDirect.app',help='Use resources and engine from this app, including an extracted release ZIP')
parser.add_argument('--live',action='store_true')
parser.add_argument('--stay-open',action='store_true')
parser.add_argument('--privacy-only',action='store_true',help='Run only the native tutor-sharing and policy checks')
parser.add_argument('--groups-only',action='store_true',help='Run only the native group-selection checks')
parser.add_argument('--excel-only',action='store_true',help='Run only native Excel import and password checks')
parser.add_argument('--forms-only',action='store_true',help='Run only native prepared-score and embedded-form checks')
args=parser.parse_args()
subprocess.run(['bash',str(root/'Scripts/test-form-descriptors.sh')],cwd=root,check=True)
source_app=args.app.resolve()
source_info=plistlib.loads((source_app/'Contents/Info.plist').read_bytes())
assert (source_app/'Contents/MacOS'/source_info['CFBundleExecutable']).is_file(), source_app
out=root/'.build/feedback-native';out.mkdir(exist_ok=True)
# A formatted file-backed source exercises insertion preflight through the real native host.
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.formula import ArrayFormula
fixture=openpyxl.Workbook(); sheet=fixture.active; sheet.title='Data'
sheet.append(['dose','double']); sheet.append([2,'=A2*2']); sheet.merge_cells('E3:F3'); sheet['E3']='keep'
sheet.add_table(Table(displayName='Doses',ref='A1:B2'))
sheet['I3']=ArrayFormula(ref='I3:J3',text='=A2:B2*2')
sheet.column_dimensions['A'].width=27; sheet['A2'].number_format='0.00'
fixture.defined_names.add(DefinedName('Dose',attr_text='Data!$A$2'))
fixture.save(out/'column-insert.xlsx')
subprocess.run([str(root/'Grid/node_modules/esbuild/bin/esbuild'),'Tests/feedback-renderer.mjs','--bundle','--format=iife','--outfile=.build/feedback-renderer.js'],cwd=root,check=True)
(out/'Viewer.swift').write_text((root/'Sources/main.swift').read_text().split('MainActor.assumeIsolated {')[0])
helpers=[]
for name in ['learning-workspace','provider-learning','report-export','group-identifier','write-back','excel-import','form-contract','form-controls']:
 p=out/(name+'.swift');p.write_text((root/('Tests/'+name+'-driver.swift')).read_text().replace('@main struct','struct'));helpers.append(str(p))
helpers.append(str(root/'Tests/report-transfer-checks.swift'))
helpers.append(str(root/'Tests/report-deletion-checks.swift'))
app=out/'StatsDirect Feedback Check.app';(app/'Contents/MacOS').mkdir(parents=True,exist_ok=True)
info=plistlib.loads((root/'Info.plist').read_bytes());info['CFBundleIdentifier']='com.statsdirect.feedback.'+uuid.uuid4().hex;info['CFBundleExecutable']='FeedbackCheck';(app/'Contents/Info.plist').write_bytes(plistlib.dumps(info))
for f in ['Resources','Frameworks']:
 p=app/'Contents'/f
 if p.is_symlink():p.unlink()
 p.symlink_to((source_app/'Contents'/f).resolve())
flags=subprocess.check_output(['bash',str(root/'Scripts/native-build-flags.sh'),'--swift'],text=True).splitlines()
cmd=['swiftc',*flags,'-emit-executable','-module-name','StatsDirectFeedback','-emit-module-path',str(out/'StatsDirectFeedback.swiftmodule'),'-parse-as-library','-target','arm64-apple-macosx14.0','-module-cache-path',str(root/'.build/swift-cache')]+[str(x) for x in (root/'Sources').glob('*.swift') if x.name!='main.swift']+[str(out/'Viewer.swift'),str(root/'Tests/beta-feedback-driver.swift')]+helpers+['-o',str(app/'Contents/MacOS/FeedbackCheck'),'-framework','Cocoa','-framework','WebKit','-framework','PDFKit','-framework','Security']
subprocess.run(cmd,cwd=root,check=True)
print('Testing app resources and engine:',source_app,flush=True)
subprocess.run([str(app/'Contents/MacOS/FeedbackCheck'),str(root/'.build/beta-feedback-exports')]+(['--live'] if args.live else [])+(['--groups-only'] if args.groups_only else [])+(['--privacy-only'] if args.privacy_only else [])+(['--excel-only'] if args.excel_only else [])+(['--forms-only'] if args.forms_only else [])+(['--stay-open'] if args.stay_open else []),cwd=root,check=True)
