"""Exercise the native Excel bridge and actual grid model; compare XLSX independently.

Run with Python + openpyxl, supplying the native workbook driver and Node paths.
The source has implicit cell coordinates, which openpyxl's shared-formula reader
cannot handle; normalize those coordinates in memory without changing the file.
"""
import copy
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'

def independent_read(path, data_only=False):
    output = io.BytesIO()
    with ZipFile(path) as source, ZipFile(output, 'w') as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
                xml = ET.fromstring(data)
                previous_row = 0
                for row in xml.iter(NS + 'row'):
                    r = int(row.get('r', previous_row + 1)); previous_row = r
                    row.set('r', str(r)); previous_col = 0
                    for cell in row.findall(NS + 'c'):
                        if cell.get('r'):
                            _, previous_col = openpyxl.utils.cell.coordinate_to_tuple(cell.get('r'))
                        else:
                            previous_col += 1
                            cell.set('r', openpyxl.utils.get_column_letter(previous_col) + str(r))
                data = ET.tostring(xml)
            dest.writestr(item, data)
    output.seek(0)
    return openpyxl.load_workbook(output, data_only=data_only)

driver = subprocess.Popen([sys.argv[1], str(ROOT/'FullEngine/publish/StatsDirectEngine.dylib')], stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)
def request(**payload):
    driver.stdin.write(json.dumps(payload) + '\n'); driver.stdin.flush()
    return json.loads(driver.stdout.readline())

try:
    source = ROOT/'Content/Examples/test.xlsx'
    imported = request(action='open', path=str(source))
    assert 'error' not in imported, imported
    assert len(imported['sheets']) == 11 and imported['formulaCount'] == 22
    exported = json.loads(subprocess.check_output([sys.argv[2], str(ROOT/'Tests/edit-workbook.mjs')], input=json.dumps(imported), text=True))
    with tempfile.TemporaryDirectory() as folder:
        saved = Path(folder)/'edited.xlsx'
        result = request(action='save', id=imported['id'], path=str(saved), **exported)
        assert result.get('ok'), result
        before, after = independent_read(source), independent_read(saved)
        assert before.sheetnames == after.sheetnames
        count = formulas = 0
        for a in before:
            b = after[a.title]
            assert a.sheet_state == b.sheet_state
            for row in a:
                for cell in row:
                    if cell.value is None: continue
                    target = b[cell.coordinate]
                    expected = 332 if a.title == 'Parametric' and cell.coordinate == 'A2' else cell.value
                    assert target.value == expected, (a.title, cell.coordinate, expected, target.value)
                    assert target.data_type == cell.data_type, (a.title, cell.coordinate, 'type')
                    assert target.number_format == cell.number_format, (a.title, cell.coordinate, 'format')
                    for feature in ('font', 'fill', 'alignment', 'border'):
                        assert copy.copy(getattr(target, feature)) == copy.copy(getattr(cell, feature)), (a.title, cell.coordinate, feature, str(getattr(cell, feature)), str(getattr(target, feature)))
                    count += 1; formulas += cell.data_type == 'f'
        old_values, new_values = independent_read(source, True), independent_read(saved, True)
        for sheet in before:
            for row in sheet:
                for cell in row:
                    if cell.data_type == 'f':
                        assert old_values[sheet.title][cell.coordinate].value == new_values[sheet.title][cell.coordinate].value, (sheet.title, cell.coordinate, 'formula result')
        reopened = request(action='open', path=str(saved))
        assert len(reopened['sheets']) == 11 and reopened['formulaCount'] == 22
        print(f'PASS: all 11 sheets, {count} populated cells, {formulas} formulas, cached formula results, cell types and cell styles survive the edited round trip')
        recalculated = Path(folder)/'recalculated.xlsx'
        changed_input = request(action='save', id=imported['id'], path=str(recalculated), sheets=[{'name':'Graphics','cells':[{'col':19,'row':1,'text':'100','kind':'number'}]}])
        assert changed_input.get('ok'), changed_input
        calc_values = independent_read(recalculated, True)['Graphics']
        assert calc_values['V2'].value == 100 - calc_values['U2'].value
        assert calc_values['W2'].value == 100 + calc_values['U2'].value
        assert independent_read(recalculated)['Graphics']['V2'].value == before['Graphics']['V2'].value
        print('PASS: editing a formula input refreshes saved arithmetic results while retaining the formula expressions')
        types_file = Path(folder)/'types.xlsx'
        values = [('00123', 'text'), ('-1.25e-8', 'number'), ('TRUE', 'boolean'), ('2024-02-29 12:34:56', 'datetime'), ('=1+1', 'text'), ('', 'blank')]
        result = request(action='save', path=str(types_file), sheets=[{'name':'Types', 'cells':[{'col':i,'row':0,'text':v,'kind':k} for i,(v,k) in enumerate(values)]}])
        assert result.get('ok'), result
        typed = independent_read(types_file)['Types']
        assert typed['A1'].value == '00123' and typed['A1'].data_type == 's'
        assert typed['B1'].value == -1.25e-8 and typed['B1'].data_type == 'n'
        assert typed['C1'].value is True
        assert str(typed['D1'].value) == '2024-02-29 12:34:56'
        assert typed['E1'].value == '=1+1' and typed['E1'].data_type == 's'
        assert typed['F1'].value is None
        print('PASS: number, text identifier, boolean, date, literal formula-looking text and blank export correctly')
        formula = next((s, c) for s in imported['sheets'] for c in s['cells'] if c['formula'])
        protected = request(action='save', id=imported['id'], path=str(saved), sheets=[{'name':formula[0]['name'],'cells':[dict(formula[1], text='0',kind='number')]}])
        assert 'read-only' in protected.get('error','')
        assert independent_read(saved)['Parametric']['A2'].value == 332
        bad = Path(folder)/'bad.xlsx'; bad.write_text('not a workbook')
        assert 'error' in request(action='open', path=str(bad))
        assert request(action='close', id=imported['id']).get('ok')
        assert 'error' in request(action='save', id=imported['id'], path=str(saved), **exported)
        print('PASS: rejected formula edits leave the destination intact; corrupt files and closed handles fail clearly')
finally:
    driver.stdin.close(); driver.wait(timeout=15)
