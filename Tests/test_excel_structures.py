"""Advanced structural edits through the production C ABI, preserving package features.

Native Excel-authored pivot/slicer fixture plus deterministic OOXML extension fixtures.
Run with workbook-driver. Outputs remain in .build/excel-structures for native Excel QA.
"""
import copy,hashlib,json,subprocess,sys,zipfile
from pathlib import Path
from xml.etree import ElementTree as E
import openpyxl
from openpyxl.worksheet.table import Table, TableColumn, TableFormula
from openpyxl.styles import Font,PatternFill
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.comments import Comment
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'.build/excel-structures'; OUT.mkdir(parents=True,exist_ok=True)
S='http://schemas.openxmlformats.org/spreadsheetml/2006/main'; NS='{'+S+'}'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'; P='http://schemas.openxmlformats.org/package/2006/relationships'
X14='http://schemas.microsoft.com/office/spreadsheetml/2009/9/main'; XM='http://schemas.microsoft.com/office/excel/2006/main'
XDR='http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'
C='http://schemas.openxmlformats.org/package/2006/content-types'
proc=subprocess.Popen([sys.argv[1],str(ROOT/'FullEngine/publish/StatsDirectEngine.dylib')],stdin=subprocess.PIPE,stdout=subprocess.PIPE,text=True)
def request(**kw):
    proc.stdin.write(json.dumps(kw)+'\n');proc.stdin.flush();return json.loads(proc.stdout.readline())
def cell(c,r,text,kind='text'): return dict(col=c,row=r,text=str(text),kind=kind)
def xml(path,part):
    with zipfile.ZipFile(path) as z:return E.fromstring(z.read(part))
def data(path,part):
    with zipfile.ZipFile(path) as z:return z.read(part)
def rewrite(source,dest,replacements={},parts={},relationships={},types={}):
    with zipfile.ZipFile(source) as z: content={n:z.read(n) for n in z.namelist()}
    for name,fn in replacements.items():content[name]=fn(content[name].decode()).encode()
    for path,new in relationships.items():
        original=content.get(path,('<Relationships xmlns="'+P+'"></Relationships>').encode()).decode()
        content[path]=original.replace('</Relationships>',''.join(f'<Relationship Id="{i}" Type="{t}" Target="{p}"/>' for i,t,p in new)+'</Relationships>').encode()
    for name,body in parts.items():content[name]=body.encode() if isinstance(body,str) else body
    content['[Content_Types].xml']=content['[Content_Types].xml'].replace(b'</Types>',(''.join(f'<Override PartName="/{p}" ContentType="{t}"/>' for p,t in types.items())+'</Types>').encode())
    with zipfile.ZipFile(dest,'w',zipfile.ZIP_DEFLATED) as z:
        for n,v in content.items():z.writestr(n,v)
    return dest
def insert(path,ops=None,stream=False,changes=None,expect=None,label=''):
    opened=request(action='open',path=str(path));assert 'error' not in opened,opened
    ops=ops or [dict(name='Data',inserts=[dict(col=1,count=1)])]
    before=hashlib.sha256(path.read_bytes()).digest()
    args=dict(id=opened['id'],inserts=ops)
    pre=request(action='validateInserts',**args)
    dest=OUT/(path.stem+label+('-stream' if stream else '-small')+'.xlsx');dest.write_bytes(b'previous destination')
    result=request(action='save',path=str(dest),stream=stream,sheets=changes or [dict(name=ops[0]['name'],cells=[])],**args)
    if expect:
        assert expect in pre.get('error','') and expect in result.get('error',''),(pre,result)
        assert dest.read_bytes()==b'previous destination'
    else:
        assert pre.get('ok') and result.get('ok'),(pre,result)
        reopened=request(action='open',path=str(dest));assert 'error' not in reopened,reopened
        request(action='close',id=reopened['id'])
    assert hashlib.sha256(path.read_bytes()).digest()==before
    assert not list(OUT.glob('*.tmp*'))
    request(action='close',id=opened['id']);return dest

def basic(name):
    w=openpyxl.Workbook();w._fonts[0]=Font(name='Calibri',sz=11);d=w.active;d.title='Data'
    for row in [['Group','Count','Double'],['A',10,'=B2*2'],['B',20,'=B3*2'],['A',30,'=B4*2']]:d.append(row)
    w.create_sheet('Other')['A1']='=Data!B2';path=OUT/name
    return w,d,path
try:
    w,d,source=basic('tables.xlsx')
    d.merge_cells('A8:C8');d['A8']='Merged heading';d['A8'].font=Font(bold=True)
    d['B1'].fill=PatternFill('solid',fgColor='FFFF00')
    d.add_table(Table(displayName='Trial',ref='A1:C4',tableColumns=[TableColumn(id=10,name='Group'),TableColumn(id=40,name='Count'),TableColumn(id=60,name='Double',calculatedColumnFormula=TableFormula(attr_text='Trial[[#This Row],[Count]]*2'))]))
    w['Other']['A2']='=SUM(Trial[Count])';w.save(source)
    for stream in (False,True):
        out=insert(source,stream=stream,changes=[dict(name='Data',cells=[cell(1,0,'Derived'),cell(1,1,99,'number')])])
        t=xml(out,'xl/tables/table1.xml');fields=t.find(NS+'tableColumns')
        assert t.get('ref')=='A1:D4' and fields.get('count')=='4'
        assert [(x.get('id'),x.get('name')) for x in fields]==[('10','Group'),('61','Derived'),('40','Count'),('60','Double')]
        x=openpyxl.load_workbook(out);assert x['Data']['B1'].value=='Derived' and x['Data']['C2'].value==10 and x['Data']['D2'].value=='=C2*2'
        assert str(x['Data'].merged_cells)=='A8:D8' and x['Other']['A2'].value=='=SUM(Trial[Count])'
        assert fields[-1].find(NS+'calculatedColumnFormula').text=='Trial[[#This Row],[Count]]*2'
        assert data(source,'xl/styles.xml')==data(out,'xl/styles.xml')
        # Multiple insertions use the current coordinates; IDs stay stable and new
        # duplicate/numeric/blank headings are normalised consistently in XML and cells.
        out=insert(source,ops=[dict(name='Data',inserts=[dict(col=1,count=2),dict(col=3,count=1),dict(col=0,count=1)])],stream=stream,label='-multiple',changes=[dict(name='Data',cells=[cell(2,0,'Count'),cell(3,0,'Count'),cell(4,0,'')])])
        x=openpyxl.load_workbook(out);t=x['Data'].tables['Trial'];assert t.ref=='B1:G4'
        assert [f.name for f in t.tableColumns]==['Group','Count2','Count3','Column1','Count','Double'],[f.name for f in t.tableColumns]
        assert [x['Data'].cell(1,c).value for c in range(2,8)]==[f.name for f in t.tableColumns]
    print('PASS: table expansion, calculated/structured references, stable column IDs, unique headings, merges, sequential edits and original styles on both save paths')

    w,d,source=basic('styles.xlsx');d['A2'].number_format='0.00%';d['A2'].font=Font(bold=True,color='FF0000');d.column_dimensions['A'].width=25
    w.save(source)
    for stream in (False,True):
        out=insert(source,stream=stream,changes=[dict(name='Data',cells=[cell(1,1,.25,'number')])]);x=openpyxl.load_workbook(out)['Data']
        assert x['B2'].value==.25 and x['B2'].number_format=='0.00%' and x['B2'].font.bold
        assert x.column_dimensions['B'].width==25
    print('PASS: inserted cells and columns inherit adjacent number formats, fonts and widths')

    # Genuine Microsoft Excel pivot/slicer package. Never rewrite record indices,
    # slicer/cache relationships or connection payloads when moving worksheet cells.
    source=ROOT/'Tests/Fixtures/Excel/structural-pivot-slicer.xlsx'
    for stream in (False,True):
        out=insert(source,ops=[dict(name='Data',inserts=[dict(col=0,count=2)]),dict(name='Sheet1',inserts=[dict(col=0,count=1)])],stream=stream)
        assert xml(out,'xl/pivotCache/pivotCacheDefinition1.xml').find('.//'+NS+'worksheetSource').get('ref')=='C1:E5'
        assert xml(out,'xl/pivotTables/pivotTable1.xml').find(NS+'location').get('ref')=='B3:C6'
        assert [e.text for e in xml(out,'xl/drawings/drawing1.xml').findall('.//{'+XDR+'}col')]==['17','19']
        with zipfile.ZipFile(source) as z:
            for n in z.namelist():
                if n.startswith(('xl/slicers/','xl/slicerCaches/','xl/pivotCache/pivotCacheRecords')) or n.endswith('.rels') and n!='xl/_rels/workbook.xml.rels':assert z.read(n)==data(out,n),n
        # Editing another sheet must not be blocked by the pivot. Growing its input
        # range retains the cache snapshot, exactly as an unrefreshed Excel pivot does.
        out=insert(source,ops=[dict(name='Data',inserts=[dict(col=1,count=1)])],stream=stream,label='-source',changes=[dict(name='Data',cells=[cell(1,0,'Extra'),cell(1,1,1,'number')])])
        assert xml(out,'xl/pivotCache/pivotCacheDefinition1.xml').find('.//'+NS+'worksheetSource').get('ref')=='A1:D5'
        insert(source,ops=[dict(name='Sheet1',inserts=[dict(col=1,count=1)])],stream=stream,label='-split',expect='split a PivotTable report')
    print('PASS: native pivot sources, whole reports and slicer anchors move; cache records/field identities stay intact; only splitting a pivot report is refused')

    w,d,base=basic('base.xlsx');w.save(base)
    ext=f'''<extLst><ext uri="{{05C60535-1F16-4fd2-B633-F4F36F0B64E0}}" xmlns:x14="{X14}" xmlns:xm="{XM}"><x14:sparklineGroups><x14:sparklineGroup displayEmptyCellsAs="gap"><x14:colorSeries rgb="FF376092"/><x14:sparklines><x14:sparkline><xm:f>Data!B2:C2</xm:f><xm:sqref>E2</xm:sqref></x14:sparkline></x14:sparklines></x14:sparklineGroup></x14:sparklineGroups></ext>
    <ext uri="{{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}}" xmlns:x14="{X14}" xmlns:xm="{XM}"><x14:dataValidations count="1"><x14:dataValidation type="whole" operator="between"><x14:formula1><xm:f>Data!B2</xm:f></x14:formula1><x14:formula2><xm:f>Data!C2</xm:f></x14:formula2><xm:sqref>B8:C9</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>'''
    source=rewrite(base,OUT/'extensions.xlsx',replacements={'xl/worksheets/sheet1.xml':lambda x:x.replace('</worksheet>',ext+'</worksheet>')})
    for stream in (False,True):
        out=insert(source,stream=stream)
        x=xml(out,'xl/worksheets/sheet1.xml');assert [e.text for e in x.findall('.//{'+XM+'}f')]==['Data!C2:D2','Data!C2','Data!D2']
        assert [e.text for e in x.findall('.//{'+XM+'}sqref')]==['F2','C8:D9']
        assert len(x.findall('.//{'+X14+'}formula1/{'+XM+'}f'))==1
    print('PASS: sparklines and extended validation ranges/formulas relocate once with extension schema retained')

    # Features stored in the worksheet (including cross-sheet consolidation and
    # conditional-format thresholds) must move even in preserved custom views.
    features=f'''<customSheetViews><customSheetView guid="{{33333333-3333-3333-3333-333333333333}}" topLeftCell="B2"><selection activeCell="B2" sqref="B2:C3"/></customSheetView></customSheetViews>
    <conditionalFormatting sqref="B2:B4"><cfRule type="dataBar" priority="1"><dataBar><cfvo type="formula" val="B2"/><cfvo type="formula" val="C2"/><color rgb="FF0000FF"/></dataBar></cfRule></conditionalFormatting>
    <dataConsolidate><dataRefs count="2"><dataRef sheet="Data" ref="B2:C4"/><dataRef sheet="Other" ref="A1:B3"/></dataRefs></dataConsolidate><cellWatches><cellWatch r="B2"/></cellWatches>'''
    def view_parts(text):
        root=E.fromstring(text); nodes=E.fromstring('<root xmlns="'+S+'">'+features+'</root>')
        # Keep schema order: consolidation, views and conditional rules before
        # margins; then breaks, watches and publishing items.
        margin=root.find(NS+'pageMargins'); at=list(root).index(margin)
        for node in [nodes[2],nodes[0],nodes[1]]:root.insert(at,node);at+=1
        root.append(nodes[3])
        breaks=E.Element(NS+'rowBreaks',count='2',manualBreakCount='2')
        E.SubElement(breaks,NS+'brk',id='3',min='0',max='16383',man='1')
        E.SubElement(breaks,NS+'brk',id='8',min='1',max='3',man='1')
        root.insert(list(root).index(margin)+1,breaks)
        publish=E.SubElement(root,NS+'webPublishItems',count='1')
        E.SubElement(publish,NS+'webPublishItem',id='1',divId='Table',sourceType='range',sourceRef='B2:C4',destinationFile='qa.html',autoRepublish='0')
        return E.tostring(root,encoding='unicode')
    viewbook='<customWorkbookViews><customWorkbookView name="QA" guid="{33333333-3333-3333-3333-333333333333}" activeSheetId="1" windowWidth="12000" windowHeight="8000"/></customWorkbookViews>'
    source=rewrite(base,OUT/'views.xlsx',replacements={'xl/worksheets/sheet1.xml':view_parts,'xl/workbook.xml':lambda x:x.replace('</workbook>',viewbook+'</workbook>')})
    for stream in (False,True):
        out=insert(source,stream=stream);x=xml(out,'xl/worksheets/sheet1.xml')
        assert x.find('.//'+NS+'customSheetView').get('topLeftCell')=='C2'
        assert x.find('.//'+NS+'selection[@activeCell="C2"]').get('sqref')=='C2:D3'
        assert [e.get('val') for e in x.findall('.//'+NS+'cfvo')]==['C2','D2']
        assert [e.get('ref') for e in x.findall('.//'+NS+'dataRef')]==['C2:D4','A1:B3']
        assert x.find('.//'+NS+'cellWatch').get('r')=='C2'
        assert [(e.get('min'),e.get('max')) for e in x.findall('.//'+NS+'brk')]==[('0','16383'),('2','4')]
        assert x.find('.//'+NS+'webPublishItem').get('sourceRef')=='C2:D4'
    first=insert(source,ops=[dict(name='Data',inserts=[dict(col=0,count=1)])],label='-first')
    assert xml(first,'xl/worksheets/sheet1.xml').find('.//'+NS+'brk').get('min')=='0'
    print('PASS: custom sheet views, consolidation sources, cell watches and formula-based conditional thresholds follow their owning sheets')

    # Dynamic-array/rich-value metadata uses cell cm/vm indices, not coordinates.
    metadata=f'<metadata xmlns="{S}"><metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" coerce="1" cellMeta="1"/></metadataTypes><futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}}"><xda:dynamicArrayProperties xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray" fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata><cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>'
    source=rewrite(base,OUT/'metadata.xlsx',replacements={'xl/worksheets/sheet1.xml':lambda x:x.replace('r="C2"','r="C2" cm="1"',1)},parts={'xl/metadata.xml':metadata},relationships={'xl/_rels/workbook.xml.rels':[('rMeta',R+'/sheetMetadata','metadata.xml')]},types={'xl/metadata.xml':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml'})
    for stream in (False,True):
        out=insert(source,stream=stream);assert data(out,'xl/metadata.xml')==data(source,'xl/metadata.xml')
        assert xml(out,'xl/worksheets/sheet1.xml').find('.//'+NS+'c[@r="D2"]').get('cm')=='1'
    print('PASS: dynamic-array metadata and per-cell metadata indices survive structural edits')

    # Native comments include the legacy fallback, author/person and drawing parts
    # that Excel needs to retain a modern thread on a subsequent save.
    source=ROOT/'Tests/Fixtures/Excel/structural-comments.xlsx'
    original=xml(source,'xl/threadedComments/threadedComment1.xml')[0]
    for stream in (False,True):
        out=insert(source,ops=[dict(name='Data',inserts=[dict(col=0,count=1)])],stream=stream)
        c=xml(out,'xl/threadedComments/threadedComment1.xml')[0]
        assert c.get('ref')=='B2' and c[0].text==original[0].text and c.get('id')==original.get('id')
        assert c.get('personId')==original.get('personId')
        assert xml(out,'xl/comments1.xml').find('.//'+NS+'comment').get('ref')=='B2'
    print('PASS: native threaded comments retain text, author/thread identity, legacy fallback and moved cell references')

    w,d,source=basic('whatif.xlsx');d['E2']=DataTableFormula(ref='E2:F4',dt2D=True,r1='B2',r2='C2');w.save(source)
    for stream in (False,True):
        out=insert(source,stream=stream);f=xml(out,'xl/worksheets/sheet1.xml').find('.//'+NS+'c[@r="F2"]/'+NS+'f')
        assert (f.get('ref'),f.get('r1'),f.get('r2'),f.get('t'))==('F2:G4','C2','D2','dataTable')
        insert(source,ops=[dict(name='Data',inserts=[dict(col=5,count=1)])],stream=stream,label='-split',expect='split a what-if data table')
    print('PASS: complete what-if tables move with both input references; invalid partial-table edits remain atomic')

    # Form controls: native VML anchor, linked cell and list source plus the Office
    # 2010 ctrlProp mirror and worksheet alternate-content anchor must agree.
    w,d,notes=basic('control-base.xlsx');d['B2'].comment=Comment('Fixture note','StatsDirect QA');w.save(notes)
    vx='urn:schemas-microsoft-com:office:excel'
    with zipfile.ZipFile(notes) as z:vml=next(n for n in z.namelist() if n.endswith('.vml'))
    def control_vml(text):
        root=E.fromstring(text);v='urn:schemas-microsoft-com:vml';o='urn:schemas-microsoft-com:office:office'
        shape=copy.deepcopy(root.find('{'+v+'}shape'));root.append(shape);shape.set('id','_x0000_s1027')
        shape.set('type','#_x0000_t201');shape.set('style','position:absolute;left:0;text-align:left;margin-left:59.25pt;margin-top:15pt;width:144pt;height:24pt;z-index:2')
        shape.set('fillcolor','window [65]');shape.set('strokecolor','windowText [64]')
        kind=copy.deepcopy(root.find('{'+v+'}shapetype'));kind.set('id','_x0000_t201');kind.set('{'+o+'}spt','201');root.insert(2,kind)
        client=shape.find('{'+vx+'}ClientData');client.set('ObjectType','Drop')
        for tag in ['MoveWithCells','SizeWithCells','Anchor','Row','Column']:
            for c in list(client.findall('{'+vx+'}'+tag)):client.remove(c)
        for tag,body in [('MoveWithCells',''),('Anchor','1, 0, 1, 0, 3, 0, 4, 0'),('FmlaLink','Data!$B$2'),('FmlaRange','Other!$A$1:Data!$B$4')]:
            E.SubElement(client,'{'+vx+'}'+tag).text=body
        # Use a normal single-sheet list source; cross-sheet colon syntax is invalid.
        client.find('{'+vx+'}FmlaRange').text='Data!$B$2:$B$4'
        return E.tostring(root,encoding='unicode')
    ctrl=f'<formControlPr xmlns="{X14}" objectType="Drop" fmlaLink="Data!$B$2" fmlaRange="Data!$B$2:$B$4"/>'
    control=f'<controls xmlns:r="{R}"><control shapeId="1027" r:id="rControl" name="Test dropdown"><controlPr linkedCell="Data!$B$2" listFillRange="Data!$B$2:$B$4"><anchor moveWithCells="1" xmlns:xdr="{XDR}"><from><xdr:col>1</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>1</xdr:row><xdr:rowOff>0</xdr:rowOff></from><to><xdr:col>3</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>4</xdr:row><xdr:rowOff>0</xdr:rowOff></to></anchor></controlPr></control></controls>'
    source=rewrite(notes,OUT/'controls.xlsx',replacements={vml:control_vml,'xl/worksheets/sheet1.xml':lambda x:x.replace('</worksheet>',control+'</worksheet>')},parts={'xl/ctrlProps/ctrlProp1.xml':ctrl},relationships={'xl/worksheets/_rels/sheet1.xml.rels':[('rControl','http://schemas.openxmlformats.org/officeDocument/2006/relationships/ctrlProp','../ctrlProps/ctrlProp1.xml')]},types={'xl/ctrlProps/ctrlProp1.xml':'application/vnd.ms-excel.controlproperties+xml'})
    for stream in (False,True):
        out=insert(source,stream=stream);v=xml(out,vml)
        assert v.find('.//{'+vx+'}FmlaLink').text=='Data!$C$2' and v.find('.//{'+vx+'}FmlaRange').text=='Data!$C$2:$C$4'
        assert [int(n) for n in v.find('.//{'+vx+'}ClientData[@ObjectType="Drop"]/{'+vx+'}Anchor').text.split(',')][::4]==[2,4]
        c=xml(out,'xl/ctrlProps/ctrlProp1.xml');assert c.get('fmlaLink')=='Data!$C$2'
        assert xml(out,'xl/worksheets/sheet1.xml').find('.//'+NS+'controlPr').get('linkedCell')=='Data!$C$2'
        assert [e.text for e in xml(out,'xl/worksheets/sheet1.xml').findall('.//'+NS+'anchor//{'+XDR+'}col')]==['2','4']
    print('PASS: legacy controls keep cell/list links and both VML and modern move-with-cells anchors')

    # Connections/query IDs are stable. An inserted table field is unbound so a
    # future query refresh does not erase the user-created results column.
    query=f'<queryTable xmlns="{S}" name="Query" connectionId="1"><queryTableRefresh nextId="4"><queryTableFields count="3"><queryTableField id="1" tableColumnId="10" name="Group"/><queryTableField id="2" tableColumnId="40" name="Count"/><queryTableField id="3" tableColumnId="60" name="Double"/></queryTableFields></queryTableRefresh></queryTable>'
    connection=f'<connections xmlns="{S}"><connection id="1" name="Fixture" type="5" refreshedVersion="6"><webPr url="https://example.invalid/data"/></connection></connections>'
    source=rewrite(OUT/'tables.xlsx',OUT/'connected.xlsx',parts={'xl/queryTables/queryTable1.xml':query,'xl/connections.xml':connection},relationships={'xl/tables/_rels/table1.xml.rels':[('rQuery',R+'/queryTable','../queryTables/queryTable1.xml')],'xl/_rels/workbook.xml.rels':[('rConnection',R+'/connections','connections.xml')]},types={'xl/queryTables/queryTable1.xml':'application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml','xl/connections.xml':'application/vnd.openxmlformats-officedocument.spreadsheetml.connections+xml'})
    for stream in (False,True):
        out=insert(source,stream=stream);q=xml(out,'xl/queryTables/queryTable1.xml');f=q.find('.//'+NS+'queryTableFields')
        assert f.get('count')=='4' and f[1].get('dataBound')=='0' and f[1].get('tableColumnId')=='61'
        assert data(out,'xl/connections.xml')==data(source,'xl/connections.xml')
        assert xml(out,'xl/tables/table1.xml').find(NS+'tableColumns')[1].get('queryTableFieldId')==f[1].get('id')
    print('PASS: connected tables add unbound query fields and preserve external connection definitions')
    params=f'<parameters count="2"><parameter name="Input" sqlType="4" parameterType="cell" cell="Data!$B$2"/><parameter name="Other" sqlType="4" parameterType="cell" cell="Other!$B$2"/></parameters>'
    source=rewrite(source,OUT/'parameters.xlsx',replacements={'xl/connections.xml':lambda x:x.replace('</connection>',params+'</connection>')})
    for stream in (False,True):
        out=insert(source,stream=stream)
        assert [p.get('cell') for p in xml(out,'xl/connections.xml').findall('.//'+NS+'parameter')]==['Data!$C$2','Other!$B$2']
        assert xml(out,'xl/connections.xml').find('.//'+NS+'webPr').get('url')=='https://example.invalid/data'
    print('PASS: cell-based connection parameters move by referenced sheet; query URLs and IDs remain intact')

    # IDs occupy unsigned 32-bit space, not necessarily small consecutive ordinals.
    source=rewrite(OUT/'tables.xlsx',OUT/'high-ids.xlsx',replacements={'xl/tables/table1.xml':lambda x:x.replace('id="60"','id="4294967295"')})
    out=insert(source,stream=True);fields=xml(out,'xl/tables/table1.xml').find(NS+'tableColumns')
    assert [f.get('id') for f in fields]==['10','1','40','4294967295']
    # Header cells can be absent from a sparse source even though its table defines
    # names. Inserting into an empty sheetData must still create the new heading.
    source=rewrite(OUT/'tables.xlsx',OUT/'empty-headers.xlsx',replacements={'xl/worksheets/sheet1.xml':lambda x:x[:x.index('<sheetData>')]+'<sheetData/>'+x[x.index('</sheetData>')+len('</sheetData>'):]})
    out=insert(source,stream=True);assert openpyxl.load_workbook(out)['Data']['B1'].value=='Column1'
    print('PASS: high field IDs allocate without overflow; sparse/empty table rows receive valid new headers')
finally:
    proc.stdin.close();proc.wait()
